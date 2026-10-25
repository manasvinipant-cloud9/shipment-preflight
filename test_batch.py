"""Batch mode: the official Bookings + Items workbook layout (template title row, header row, example row)."""
import io
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE / "samples"))
from build_samples import B_HEAD, EXPECTED, I_HEAD  # noqa: E402

from preflight.sheets import SheetError, confirmations_text, findings_csv, run_batch  # noqa: E402

TODAY = date(2026, 10, 25)
XLSX = (HERE / "samples" / "test_bookings_items.xlsx").read_bytes()


def by_id(results):
    return {r.booking_id: r for r in results}


def workbook(bookings, items, extra=None):
    wb = Workbook()
    ws = wb.active; ws.title = "Bookings"
    ws.append(["Bookings — one row per record. Row 3 is an example (grey); enter data from row 4."]); ws.append(B_HEAD)
    for r in bookings: ws.append(r)
    ws = wb.create_sheet("Items")
    ws.append(["Items — one row per record."]); ws.append(I_HEAD)
    for r in items: ws.append(r)
    for title, rows in (extra or {}).items():
        ws = wb.create_sheet(title)
        for r in rows: ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return [("t.xlsx", buf.getvalue())]


BK = ["BK-X", "Seller", "Leeds", "United Kingdom", "Paris", "France", "Business address", "10-11-2026", "Balanced", 1, 10,
      100, "GBP", 60, "No"]
IT = ["BK-X", "S1", "Mug", 691200, 10, 10, 1, 10, 1, "India"]


def test_sample_workbook_hits_exactly_the_planted_rules():
    results, notes = run_batch([("t.xlsx", XLSX)], TODAY)
    assert notes == [] and len(results) == 10
    for r in results:
        got = {f.rule_id for f in r.report.errors + r.report.warnings}
        assert got == EXPECTED[r.booking_id], (r.booking_id, got)
        assert r.notes == [], (r.booking_id, r.notes)


def test_sample_recommendations_follow_priority_and_weight():
    r = by_id(run_batch([("t.xlsx", XLSX)], TODAY)[0])
    rec = lambda b: r[b].report.recommendation
    assert (rec("BK-001").lane.name, rec("BK-001").cost_gbp) == ("Carrier A", 540)      # Cheapest, 3 pallets x 180
    assert rec("BK-002").lane.name == "Carrier B"                                       # Balanced
    assert rec("BK-004").lane.name == "Carrier C"                                       # Fastest
    assert rec("BK-008").lane.name == "Carrier C"                                       # Most reliable
    assert rec("BK-006").lane.name == "Carrier B" and r["BK-006"].report.provisional    # 650 kg/pallet: A can't, B can
    assert rec("BK-010") is None                                                        # 800 kg/pallet: nothing takes it
    assert r["BK-001"].report.status == "READY" and r["BK-008"].report.status == "REVIEW"


def test_csv_pair_matches_workbook():
    a, _ = run_batch([("t.xlsx", XLSX)], TODAY)
    b, _ = run_batch([("bookings.csv", (HERE / "samples/bookings.csv").read_bytes()),
                      ("items.csv", (HERE / "samples/items.csv").read_bytes())], TODAY)
    assert [(x.booking_id, x.report.status) for x in a] == [(x.booking_id, x.report.status) for x in b]


def test_template_titles_dates_and_numbers_are_read():
    r = by_id(run_batch([("t.xlsx", XLSX)], TODAY)[0])
    s = r["BK-002"].shipment                       # datetime cell
    assert s.ready_date == date(2026, 11, 9) and s.destination_type == "Business address"
    assert r["BK-001"].shipment.ready_date == date(2026, 11, 5)   # '05-11-2026' text, DD-MM-YYYY
    assert r["BK-001"].shipment.items[0].hs_code == "691200"


def test_data_quality_notes():
    bad = list(BK); bad[8] = None; bad[12] = "USD"; bad[9] = None
    res, notes = run_batch(workbook([bad], [IT, ["BK-ORPHAN"] + IT[1:]]), TODAY)
    n = " ".join(res[0].notes)
    assert "assumed Balanced" in n and "USD" in n and "pallets" in n
    assert any("BK-ORPHAN" in x for x in notes)
    res, notes = run_batch(workbook([BK, BK], [IT]), TODAY)
    assert any("Duplicate" in x for x in notes)
    res, _ = run_batch(workbook([BK], []), TODAY)
    assert any("No items" in x for x in res[0].notes)


def test_blank_hs_and_origin_are_allowed_but_flagged():
    it = list(IT); it[3] = None; it[9] = None
    r = run_batch(workbook([BK], [it]), TODAY)[0][0]
    assert {f.rule_id for f in r.report.errors + r.report.warnings} >= {"R-01", "R-09"}
    assert not any("hs_code" in n or "country_of_origin" in n for n in r.notes)    # legitimately blank, not "missing data"


def test_rate_card_override_from_carrier_options_sheet():
    card = [["FIXED - judges will use exactly this rate card. Do not edit."],
            ["Carrier", "Service", "Origin country", "Destination country", "Price per pallet", "Currency",
             "Transit time (days)", "On-time rate (%)", "Max weight per pallet (kg)"],
            ["Carrier Z", "Road economy", "United Kingdom", "France", 50, "GBP", 4, 90, 900]]
    res, notes = run_batch(workbook([BK], [IT], {"Carrier options": card}), TODAY)
    assert res[0].report.recommendation.lane.name == "Carrier Z" and any("rate card" in n for n in notes)


def test_bad_files_give_clear_errors():
    for files in ([("x.txt", b"hi")], [("x.csv", b"a,b\n1,2\n")], [], [("x.xlsx", b"not a zip")]):
        try:
            run_batch(files, TODAY)
        except SheetError:
            continue
        except Exception as e:                       # unreadable xlsx: any error is caught by the UI, but prefer SheetError
            assert "x.xlsx" in str(files), e
            continue
        raise AssertionError(files)
    try:
        run_batch(workbook([], [IT]), TODAY)
    except SheetError as e:
        assert "no data rows" in str(e)
    else:
        raise AssertionError("empty bookings sheet")


def test_exports():
    res, _ = run_batch([("t.xlsx", XLSX)], TODAY)
    csv_text = findings_csv(res)
    assert csv_text.splitlines()[0].startswith("booking_id,seller,status") and "R-05" in csv_text and "BK-001" in csv_text
    txt = confirmations_text(res)
    assert "BK-006" in txt and "What happens next" in txt and "Why this carrier" in txt


if __name__ == "__main__":
    fails = 0
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            try:
                f(); print("PASS", n)
            except Exception as e:
                import traceback; traceback.print_exc(); print("FAIL", n, repr(e)); fails += 1
    sys.exit(fails)
